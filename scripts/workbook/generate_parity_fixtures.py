#!/usr/bin/env python3
"""Generate workbook parity fixtures for the TypeScript calculation engine.

For each scenario below, a copy of the reference workbook is filled with the
scenario's cell values, recalculated headlessly by LibreOffice, and every
formula cell's result is captured. The TypeScript tests replay the same inputs
through the domain layer and compare against these workbook-computed results.

All scenario values are SYNTHETIC TEST VALUES chosen to exercise the formulas.
They are not hospital data and must never be shown as such.

Requirements: Python 3.10+, openpyxl, LibreOffice (`soffice` on PATH).

Usage:
    python3 scripts/workbook/generate_parity_fixtures.py \
        [--workbook data/Elite_RT_Value_Calculator.xlsx] \
        [--out tests/fixtures/workbook-parity.json]

openpyxl drops cached formula results when it saves, so LibreOffice has to
compute every formula on load; the script checks that it did.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = REPO_ROOT / "data" / "Elite_RT_Value_Calculator.xlsx"
DEFAULT_OUT = REPO_ROOT / "tests" / "fixtures" / "workbook-parity.json"

# Synthetic, fully populated input set (every yellow cell on Inputs and Revenue).
FULL_INPUTS: dict[str, float] = {
    "Inputs!C7": 0.78,
    "Inputs!C8": 4200,
    "Inputs!C9": 5.5,
    "Inputs!C10": 1800,
    "Inputs!C11": 1500,
    "Inputs!C12": 250000,
    "Inputs!C14": 2500,
    "Inputs!C15": 900,
    "Inputs!C16": 1100,
    "Inputs!C18": 3000000,
    "Inputs!C19": 4500000,
    "Inputs!C20": 600000,
    "Inputs!C21": 850000,
    "Inputs!C22": 2000000,
    "Inputs!C23": 9000000,
    "Inputs!C24": 1200000,
    "Inputs!C25": 400000,
    "Inputs!C26": 750000,
    "Inputs!C28": 3000,
    "Inputs!C29": 12000,
    "Inputs!C30": 5000,
    "Inputs!C31": 1400,
    "Inputs!C32": 900,
    "Inputs!C33": 85,
    "Inputs!C34": 0.62,
    "Inputs!C36": 20000,
    "Inputs!C37": 3500,
    "Inputs!C38": 350,
    "Inputs!C39": 0.85,
    "Revenue!B39": 1200,
    "Revenue!C39": 1500,
    "Revenue!B40": 300,
    "Revenue!C40": 4000,
    "Revenue!B41": 50,
    "Revenue!C41": 20000,
    "Revenue!B45": 4800000,
    "Revenue!B46": 1200000,
    "Revenue!B47": 300000,
    "Revenue!B48": 900000,
    "Revenue!B49": 1100000,
    "Revenue!B50": 400000,
    "Revenue!B51": 250000,
    "Revenue!B52": 200000,
    "Revenue!B53": 600000,
    "Revenue!B54": 150000,
}

OPERATING_COST_CELLS = [f"Revenue!B{r}" for r in range(45, 55)]

# A value of None clears the cell (blank = unknown, distinct from 0).
SCENARIOS: list[dict] = [
    {
        "id": "as-shipped",
        "description": "Workbook exactly as delivered: only ICU beds = 50, selectors 80% / EGP 1,000 / Mid.",
        "cells": {},
    },
    {
        "id": "full-mid-80-1000",
        "description": "All inputs populated (synthetic); default selectors 80% / EGP 1,000 / Mid.",
        "cells": {**FULL_INPUTS},
    },
    {
        "id": "full-low-60-800",
        "description": "All inputs populated (synthetic); selectors 60% / EGP 800 / Low.",
        "cells": {**FULL_INPUTS, "Value Bridge!C4": 0.6, "Value Bridge!C5": 800, "Value Bridge!C6": "Low"},
    },
    {
        "id": "full-high-90-1200",
        "description": "All inputs populated (synthetic); selectors 90% / EGP 1,200 / High.",
        "cells": {**FULL_INPUTS, "Value Bridge!C4": 0.9, "Value Bridge!C5": 1200, "Value Bridge!C6": "High"},
    },
    {
        "id": "partial-components",
        "description": (
            "Partial data: ventilator days without cost; NIV pair complete but HFNC cost missing; "
            "only equipment rental; only overtime; billing without collection rate; PFT volume without price; "
            "one of ten operating-cost lines entered (workbook still reports a net value)."
        ),
        "cells": {
            "Inputs!C8": 4200,
            "Inputs!C10": 1800,
            "Inputs!C15": 900,
            "Inputs!C11": 1500,
            "Inputs!C20": 600000,
            "Inputs!C24": 1200000,
            "Inputs!C37": 3500,
            "Inputs!C38": 350,
            "Revenue!B39": 1200,
            "Revenue!B45": 4800000,
        },
    },
    {
        "id": "niv-incomplete-hfnc-complete",
        "description": "NIV days without NIV cost; HFNC pair complete. NIV contributes 0, HFNC is counted.",
        "cells": {"Inputs!C10": 1800, "Inputs!C11": 1500, "Inputs!C16": 1100},
    },
    {
        "id": "niv-cost-only",
        "description": "Only cost per NIV day entered: both NIV/HFNC pairs incomplete, so the lever is not quantified.",
        "cells": {"Inputs!C15": 900},
    },
    {
        "id": "zero-vs-blank",
        "description": (
            "Explicit zeros: ventilated days 0, oxygen spend 0, unbilled activities 0, all ten operating-cost "
            "lines 0. Zeros are data, so results are calculated as 0 rather than 'Baseline required'."
        ),
        "cells": {
            "Inputs!C8": 0,
            "Inputs!C14": 2500,
            "Inputs!C18": 0,
            "Inputs!C37": 0,
            "Inputs!C38": 350,
            "Inputs!C39": 0.85,
            **{cell: 0 for cell in OPERATING_COST_CELLS},
        },
    },
    {
        "id": "edited-assumptions",
        "description": (
            "Respiratory Gate assumptions edited: 45 beds, 31-day month, 366-day year, occupancy grid "
            "55/65/75/85%, price grid 900/1,100/1,300, savings 3/8/20%; selectors 75% / EGP 1,100 / Low "
            "(values outside the workbook dropdown lists)."
        ),
        "cells": {
            **FULL_INPUTS,
            "Inputs!C6": 45,
            "Revenue!B5": 31,
            "Revenue!B6": 366,
            "Revenue!C9": 900,
            "Revenue!D9": 1100,
            "Revenue!E9": 1300,
            "Revenue!A11": 0.55,
            "Revenue!A12": 0.65,
            "Revenue!A13": 0.75,
            "Revenue!A14": 0.85,
            "Savings Scenarios!D5": 0.03,
            "Savings Scenarios!E5": 0.08,
            "Savings Scenarios!F5": 0.2,
            "Value Bridge!C4": 0.75,
            "Value Bridge!C5": 1100,
            "Value Bridge!C6": "Low",
        },
    },
    {
        "id": "icu-beds-blank",
        "description": (
            "ICU beds cleared. The workbook treats the blank as 0 and reports ICU package revenue as a "
            "'Calculated' 0; the web app deliberately reports 'Data required' instead (see formula map)."
        ),
        "cells": {"Inputs!C6": None},
    },
]


def split_ref(ref: str) -> tuple[str, str]:
    sheet, cell = ref.rsplit("!", 1)
    return sheet, cell


def normalise(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def soffice_binary() -> str:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    sys.exit("LibreOffice (soffice) not found on PATH.")


def soffice_version(binary: str) -> str:
    out = subprocess.run([binary, "--version"], capture_output=True, text=True, check=False)
    return out.stdout.strip()


def formula_cells(workbook_path: Path) -> dict[str, list[str]]:
    wb = openpyxl.load_workbook(workbook_path, data_only=False)
    cells: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        refs = [
            c.coordinate
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("=")
        ]
        if refs:
            cells[ws.title] = refs
    return cells


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    workbook_path: Path = args.workbook.resolve()
    sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    binary = soffice_binary()
    formulas = formula_cells(workbook_path)

    with tempfile.TemporaryDirectory(prefix="rg-parity-") as tmp:
        tmp_path = Path(tmp)
        in_dir, out_dir, profile = tmp_path / "in", tmp_path / "out", tmp_path / "profile"
        in_dir.mkdir()
        out_dir.mkdir()

        for scenario in SCENARIOS:
            wb = openpyxl.load_workbook(workbook_path)
            for ref, value in scenario["cells"].items():
                sheet, cell = split_ref(ref)
                wb[sheet][cell].value = value
            wb.save(in_dir / f"{scenario['id']}.xlsx")

        cmd = [
            binary,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--norestore",
            "--calc",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            *sorted(str(p) for p in in_dir.glob("*.xlsx")),
        ]
        subprocess.run(cmd, check=True, capture_output=True, timeout=300)

        results = []
        for scenario in SCENARIOS:
            recalculated = out_dir / f"{scenario['id']}.xlsx"
            if not recalculated.exists():
                sys.exit(f"LibreOffice did not produce {recalculated.name}")
            wb = openpyxl.load_workbook(recalculated, data_only=True)

            expected: dict[str, object] = {}
            for sheet, refs in formulas.items():
                for cell in refs:
                    value = wb[sheet][cell].value
                    # A formula that returns "" is read back as None.
                    expected[f"{sheet}!{cell}"] = "" if value is None else normalise(value)

            # Guard against LibreOffice skipping recalculation: the occupied-bed
            # count is a plain product and must always be numeric.
            if not isinstance(expected["Revenue!B18"], (int, float)):
                sys.exit(f"Scenario {scenario['id']}: workbook was not recalculated")

            results.append(
                {
                    "id": scenario["id"],
                    "description": scenario["description"],
                    "cells": {ref: normalise(v) for ref, v in scenario["cells"].items()},
                    "expected": expected,
                }
            )

    fixture = {
        "_comment": (
            "Generated by scripts/workbook/generate_parity_fixtures.py. Do not edit by hand. "
            "Scenario values are SYNTHETIC TEST VALUES, not hospital data. 'cells' lists the cells changed "
            "from the delivered workbook (null = cleared); 'expected' holds every formula cell's "
            "LibreOffice-recalculated result ('' = formula returned an empty string)."
        ),
        "workbook": workbook_path.relative_to(REPO_ROOT).as_posix()
        if workbook_path.is_relative_to(REPO_ROOT)
        else workbook_path.name,
        "workbookSha256": sha256,
        "calculator": soffice_version(binary),
        "scenarios": results,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(fixture, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(results)} scenarios to {args.out}")


if __name__ == "__main__":
    main()
